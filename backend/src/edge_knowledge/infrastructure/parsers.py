from io import BytesIO
from pathlib import PurePath

from edge_knowledge.domain.models import AppError


class FileDocumentParser:
    def parse(self, filename: str, content: bytes) -> str:
        suffix = PurePath(filename).suffix.lower()
        try:
            if suffix in {".txt", ".md"}:
                try:
                    return content.decode("utf-8-sig")
                except UnicodeDecodeError:
                    return content.decode("gb18030")
            if suffix == ".pdf":
                from pypdf import PdfReader

                reader = PdfReader(BytesIO(content))
                if reader.is_encrypted:
                    raise AppError("请先解除 PDF 密码再导入。")
                return "\n\n".join(page.extract_text() or "" for page in reader.pages)
            if suffix == ".docx":
                from docx import Document

                document = Document(BytesIO(content))
                parts = [paragraph.text for paragraph in document.paragraphs]
                parts.extend(
                    " | ".join(cell.text for cell in row.cells)
                    for table in document.tables
                    for row in table.rows
                )
                return "\n".join(parts)
        except AppError:
            raise
        except Exception as exc:
            raise AppError("文件解析失败，请确认文件未损坏且格式正确。") from exc
        raise AppError("仅支持 TXT、Markdown、PDF 和 DOCX 文件。", 415, "unsupported_file")
